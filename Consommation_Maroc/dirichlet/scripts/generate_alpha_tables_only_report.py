from __future__ import annotations

import math
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook

from generate_alpha_diagnostic_report import (
    CASE_ORDER,
    DMM_CSV,
    DPMM_CSV,
    OUT_DIR,
    collect_dmm,
    collect_dpmm,
    write_csv,
)


PDF_OUT = OUT_DIR / "rapport_tableaux_alpha_DMM_DPMM.pdf"
XLSX_OUT = OUT_DIR / "tableaux_alpha_DMM_DPMM.xlsx"


def fmt(value: float | int | str | None, digits: int = 2) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    value = float(value)
    if not math.isfinite(value):
        return ""
    return f"{value:.{digits}f}"


def pct(value: float | int | None, digits: int = 2) -> str:
    if value is None:
        return ""
    return fmt(100 * float(value), digits)


def add_table_page(
    pdf: PdfPages,
    title: str,
    columns: list[str],
    rows: list[list[str]],
    font_size: float,
    scale_y: float = 1.35,
) -> None:
    fig = plt.figure(figsize=(11.69, 8.27))
    fig.patch.set_facecolor("white")
    fig.text(0.04, 0.95, title, fontsize=16, fontweight="bold", ha="left", va="top")

    ax = fig.add_axes([0.02, 0.05, 0.96, 0.86])
    ax.axis("off")
    table = ax.table(cellText=rows, colLabels=columns, loc="center", cellLoc="center")
    table.auto_set_font_size(False)
    table.set_fontsize(font_size)
    table.scale(1, scale_y)

    for (row, _col), cell in table.get_celld().items():
        cell.set_edgecolor("#B8C0CC")
        if row == 0:
            cell.set_facecolor("#E9EEF5")
            cell.set_text_props(weight="bold")
        else:
            cell.set_facecolor("#FFFFFF")

    pdf.savefig(fig, bbox_inches="tight")
    plt.close(fig)


def dmm_table_rows(rows: list[dict[str, float | str | int]]) -> list[list[str]]:
    out = []
    for row in sorted(rows, key=lambda r: int(r["K"])):
        out.append(
            [
                str(row["K"]),
                str(row["n_components"]),
                fmt(row["alpha0_min"]),
                fmt(row["alpha0_q25"]),
                fmt(row["alpha0_median"]),
                fmt(row["alpha0_q75"]),
                fmt(row["alpha0_max"]),
                pct(row["pi_min"], 2),
                pct(row["pi_median"], 2),
                pct(row["pi_max"], 2),
                fmt(row["alpha_j_min"], 3),
                fmt(row["alpha_j_max"], 2),
            ]
        )
    return out


def dpmm_table_rows(rows: list[dict[str, float | str | int]]) -> list[list[str]]:
    out = []
    for row in rows:
        out.append(
            [
                str(row["prior"]),
                str(row["prior_hyperparams"]),
                str(row["n_active_ge_1pct"]),
                str(row["n_active_ge_2pct"]),
                str(row["n_active_ge_3pct"]),
                fmt(row["alpha0_active_min"]),
                fmt(row["alpha0_active_median"]),
                fmt(row["alpha0_active_max"]),
                pct(row["weight_min"], 2),
                pct(row["weight_median"], 2),
                pct(row["weight_max"], 2),
                fmt(row["alpha_j_min"], 3),
                fmt(row["alpha_j_max"], 2),
            ]
        )
    return out


def write_workbook(dmm_rows: list[dict[str, float | str | int]], dpmm_rows: list[dict[str, float | str | int]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DMM_resume"
    write_sheet(ws, dmm_rows)

    ws = wb.create_sheet("DPMM_resume")
    write_sheet(ws, dpmm_rows)

    for case_key in CASE_ORDER:
        case_dmm = [row for row in dmm_rows if str(row["case_key"]) == case_key]
        if case_dmm:
            ws = wb.create_sheet(f"DMM_{case_key}")
            write_sheet(ws, case_dmm)
        case_dpmm = [row for row in dpmm_rows if str(row["case_key"]) == case_key]
        if case_dpmm:
            ws = wb.create_sheet(f"DPMM_{case_key}")
            write_sheet(ws, case_dpmm)

    wb.save(XLSX_OUT)


def write_sheet(ws, rows: list[dict[str, float | str | int]]) -> None:
    if not rows:
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 28)


def make_pdf(dmm_rows: list[dict[str, float | str | int]], dpmm_rows: list[dict[str, float | str | int]]) -> None:
    dmm_cols = [
        "K",
        "comp.",
        "alpha0 min",
        "alpha0 Q25",
        "alpha0 med.",
        "alpha0 Q75",
        "alpha0 max",
        "pi min %",
        "pi med. %",
        "pi max %",
        "alpha_j min",
        "alpha_j max",
    ]
    dpmm_cols = [
        "prior",
        "hyperparam.",
        "act. >=1%",
        "act. >=2%",
        "act. >=3%",
        "alpha0 min",
        "alpha0 med.",
        "alpha0 max",
        "poids min %",
        "poids med. %",
        "poids max %",
        "alpha_j min",
        "alpha_j max",
    ]

    with PdfPages(PDF_OUT) as pdf:
        for case_key in CASE_ORDER:
            case_dmm = [row for row in dmm_rows if str(row["case_key"]) == case_key]
            if case_dmm:
                add_table_page(
                    pdf,
                    f"DMM - {case_dmm[0]['case']} - alpha par K",
                    dmm_cols,
                    dmm_table_rows(case_dmm),
                    font_size=7.5,
                    scale_y=1.5,
                )

            case_dpmm = [row for row in dpmm_rows if str(row["case_key"]) == case_key]
            if case_dpmm:
                add_table_page(
                    pdf,
                    f"DPMM - {case_dpmm[0]['case']} - alpha par prior",
                    dpmm_cols,
                    dpmm_table_rows(case_dpmm),
                    font_size=6.8,
                    scale_y=1.8,
                )


def main() -> None:
    dmm_rows, _dmm_alphas = collect_dmm()
    dpmm_rows, _dpmm_alphas = collect_dpmm()
    write_csv(DMM_CSV, dmm_rows)
    write_csv(DPMM_CSV, dpmm_rows)
    write_workbook(dmm_rows, dpmm_rows)
    make_pdf(dmm_rows, dpmm_rows)
    print(PDF_OUT)
    print(XLSX_OUT)


if __name__ == "__main__":
    main()
